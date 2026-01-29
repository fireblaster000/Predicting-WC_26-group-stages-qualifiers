import os
import re
import pandas as pd
from openpyxl import load_workbook

# -----------------------------
# Paths
# -----------------------------
EXCEL_PATH = r"data\group_stage_qualification_data.xlsx"
SHEET_NAME = "Team_Group_Data"
RESULTS_CSV = r"data\results.csv"

# Choose recent window size
N_RECENT = 10

# -----------------------------
# Name normalization (expand as needed)
# -----------------------------
NAME_MAP_RESULTS = {
    "Republic of Ireland": "Ireland",
    "Korea Republic": "South Korea",
    "IR Iran": "Iran",
    "China": "China PR",
    "Serbia and Montenegro": "Serbia",
}

def norm_name(s: str) -> str:
    s = str(s).strip()
    s = NAME_MAP_RESULTS.get(s, s)
    s = re.sub(r"\s+", " ", s)
    return s

# -----------------------------
# Build team-centric results table
# -----------------------------
def load_team_results(results_csv_path: str) -> pd.DataFrame:
    res = pd.read_csv(results_csv_path)

    required = ["date", "home_team", "away_team", "home_score", "away_score"]
    missing_cols = [c for c in required if c not in res.columns]
    if missing_cols:
        raise ValueError(f"results.csv is missing required columns: {missing_cols}. "
                         f"Found columns: {list(res.columns)}")

    res["date"] = pd.to_datetime(res["date"], errors="coerce")

    home = pd.DataFrame({
        "date": res["date"],
        "team": res["home_team"].apply(norm_name),
        "opponent": res["away_team"].apply(norm_name),
        "goals_for": res["home_score"],
        "goals_against": res["away_score"],
    })

    away = pd.DataFrame({
        "date": res["date"],
        "team": res["away_team"].apply(norm_name),
        "opponent": res["home_team"].apply(norm_name),
        "goals_for": res["away_score"],
        "goals_against": res["home_score"],
    })

    team_games = pd.concat([home, away], ignore_index=True)

    team_games["goals_for"] = pd.to_numeric(team_games["goals_for"], errors="coerce")
    team_games["goals_against"] = pd.to_numeric(team_games["goals_against"], errors="coerce")
    team_games = team_games.dropna(subset=["date", "team", "goals_for", "goals_against"]).copy()

    team_games["gd"] = team_games["goals_for"] - team_games["goals_against"]
    team_games["is_win"] = (team_games["gd"] > 0).astype(int)
    team_games["is_draw"] = (team_games["gd"] == 0).astype(int)
    team_games["is_loss"] = (team_games["gd"] < 0).astype(int)

    # Sort so tail(N) is "most recent N"
    team_games = team_games.sort_values(["team", "date"]).reset_index(drop=True)
    return team_games

# -----------------------------
# Compute recent features for (team, cutoff_date)
# -----------------------------
def recent_features_for_team(team_games: pd.DataFrame, team: str, cutoff: pd.Timestamp, n_recent: int):
    games = team_games[(team_games["team"] == team) & (team_games["date"] <= cutoff)]
    if games.empty:
        return (0, None, None, None, None)

    last = games.tail(n_recent)
    k = len(last)

    win_rate = last["is_win"].mean()
    gd_pm = last["gd"].mean()
    gf_pm = last["goals_for"].mean()
    ga_pm = last["goals_against"].mean()

    return (k, float(win_rate), float(gd_pm), float(gf_pm), float(ga_pm))

# -----------------------------
# Main pipeline: compute + write in-place
# -----------------------------
def compute_recent_form_and_write():
    if not os.path.exists(RESULTS_CSV):
        raise FileNotFoundError(
            f"Missing RESULTS_CSV at {RESULTS_CSV}. Put results.csv there first."
        )

    # Read Excel to get rows we need (header is on row 3)
    df = pd.read_excel(EXCEL_PATH, sheet_name=SHEET_NAME, skiprows=2)
    req = ["tournament_year", "group_id", "team_name", "tournament_start_date"]
    df = df[df[req].notna().all(axis=1)].copy()

    df["tournament_year"] = df["tournament_year"].astype(int)
    df["group_id"] = df["group_id"].astype(str).str.strip()
    df["team_name"] = df["team_name"].astype(str).str.strip()
    df["team_name_norm"] = df["team_name"].apply(norm_name)

    df["tournament_start_date"] = pd.to_datetime(df["tournament_start_date"], errors="coerce")
    df["cutoff_date"] = df["tournament_start_date"] - pd.Timedelta(days=1)
    df["cutoff_date_str"] = df["cutoff_date"].dt.strftime("%Y-%m-%d")

    # Load results and make team-centric table
    team_games = load_team_results(RESULTS_CSV)

    # Unique cutoffs and teams we need features for
    unique_cutoffs = sorted(df["cutoff_date"].dropna().unique())
    teams_needed = sorted(df["team_name_norm"].dropna().unique())

    # Build lookup: (cutoff_date_str, team_norm) -> 5-tuple feats
    feat_lookup = {}
    for cutoff in unique_cutoffs:
        cutoff_ts = pd.to_datetime(cutoff)
        cutoff_str = cutoff_ts.strftime("%Y-%m-%d")
        for team in teams_needed:
            feat_lookup[(cutoff_str, team)] = recent_features_for_team(
                team_games, team, cutoff_ts, N_RECENT
            )

    # Columns I will write
    recent_cols = [
        "recent_window_n_matches",
        "recent_win_rate",
        "recent_goal_diff_per_match",
        "recent_goals_for_per_match",
        "recent_goals_against_per_match",
    ]

    # IMPORTANT: Avoid duplicate column names (Excel might already contain them)
    df = df.drop(columns=[c for c in recent_cols if c in df.columns], errors="ignore")

    def fetch_feats(row):
        key = (row["cutoff_date_str"], row["team_name_norm"])
        val = feat_lookup.get(key, (0, None, None, None, None))
        if val is None or len(val) != 5:
            return (0, None, None, None, None)
        return val

    feats = df.apply(fetch_feats, axis=1, result_type="expand")
    if feats.shape[1] != 5:
        raise ValueError(f"Expected feats with 5 cols, got shape {feats.shape}")

    feats.columns = recent_cols
    df = pd.concat([df, feats], axis=1)

    # Force correct dtypes (and avoid NaN->int crash)
    df["recent_window_n_matches"] = pd.to_numeric(df["recent_window_n_matches"], errors="coerce").fillna(0).astype(int)
    for c in recent_cols[1:]:
        df[c] = pd.to_numeric(df[c], errors="coerce")

    # Prepare mapping to write back: (year, group_id, team_name) -> values
    out_lookup = {}
    for r in df.itertuples(index=False):
        key = (int(r.tournament_year), str(r.group_id), str(r.team_name))
        out_lookup[key] = {
            "recent_window_n_matches": int(r.recent_window_n_matches),
            "recent_win_rate": float(r.recent_win_rate) if pd.notna(r.recent_win_rate) else None,
            "recent_goal_diff_per_match": float(r.recent_goal_diff_per_match) if pd.notna(r.recent_goal_diff_per_match) else None,
            "recent_goals_for_per_match": float(r.recent_goals_for_per_match) if pd.notna(r.recent_goals_for_per_match) else None,
            "recent_goals_against_per_match": float(r.recent_goals_against_per_match) if pd.notna(r.recent_goals_against_per_match) else None,
        }

    # Write back using openpyxl in-place
    wb = load_workbook(EXCEL_PATH)
    ws = wb[SHEET_NAME]

    header_row = 3
    header_to_col = {}
    for col in range(1, ws.max_column + 1):
        v = ws.cell(header_row, col).value
        if v:
            header_to_col[str(v).strip()] = col

    for c in recent_cols:
        if c not in header_to_col:
            raise ValueError(f"Missing column in Excel header: {c}. Add it to the template first.")

    data_start_row = 4
    rows_written = 0

    for excel_row in range(data_start_row, ws.max_row + 1):
        year = ws.cell(excel_row, header_to_col["tournament_year"]).value
        group_id = ws.cell(excel_row, header_to_col["group_id"]).value
        team = ws.cell(excel_row, header_to_col["team_name"]).value

        if year is None or group_id is None or team is None:
            continue

        key = (int(year), str(group_id).strip(), str(team).strip())
        if key not in out_lookup:
            continue

        vals = out_lookup[key]
        ws.cell(excel_row, header_to_col["recent_window_n_matches"]).value = vals["recent_window_n_matches"]
        ws.cell(excel_row, header_to_col["recent_win_rate"]).value = vals["recent_win_rate"]
        ws.cell(excel_row, header_to_col["recent_goal_diff_per_match"]).value = vals["recent_goal_diff_per_match"]
        ws.cell(excel_row, header_to_col["recent_goals_for_per_match"]).value = vals["recent_goals_for_per_match"]
        ws.cell(excel_row, header_to_col["recent_goals_against_per_match"]).value = vals["recent_goals_against_per_match"]

        rows_written += 1

    try:
        wb.save(EXCEL_PATH)
    except PermissionError:
        raise PermissionError(f"Permission denied saving '{EXCEL_PATH}'. Close the Excel file and try again.")

    print(f"Done. Wrote recent-form features (N={N_RECENT}) for {rows_written} rows into {EXCEL_PATH}")

    # Debug: teams with 0 prior matches found
    zero = df[df["recent_window_n_matches"] == 0][["tournament_year", "team_name", "team_name_norm"]].drop_duplicates()
    if len(zero) > 0:
        print("\nWARNING: Teams with 0 prior matches found (check name mapping or missing years in results.csv):")
        print(zero.to_string(index=False))


if __name__ == "__main__":
    compute_recent_form_and_write()
