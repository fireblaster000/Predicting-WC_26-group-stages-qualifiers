from openpyxl import load_workbook
import re

EXCEL_PATH = r"data\group_stage_qualification_data.xlsx"
SHEET_NAME = "Team_Group_Data"

HEADER_ROW = 3
DATA_START_ROW = 4

# -----------------------------
# TEAM → CONFEDERATION MAP
# Expand if needed
# -----------------------------
TEAM_TO_CONFED = {
    # UEFA
    "Germany": "UEFA",
    "France": "UEFA",
    "Spain": "UEFA",
    "England": "UEFA",
    "Italy": "UEFA",
    "Netherlands": "UEFA",
    "Portugal": "UEFA",
    "Belgium": "UEFA",
    "Croatia": "UEFA",
    "Switzerland": "UEFA",
    "Sweden": "UEFA",
    "Denmark": "UEFA",
    "Poland": "UEFA",
    "Serbia": "UEFA",
    "Yugoslavia": "UEFA",
    "Czech Republic": "UEFA",
    "Slovakia": "UEFA",
    "Russia": "UEFA",
    "Austria": "UEFA",
    "Bosnia and Herzegovina": "UEFA",
    "Bulgaria": "UEFA",
    "Greece": "UEFA",
    "Iceland": "UEFA",
    "Norway": "UEFA",
    "Republic of Ireland": "UEFA",
    "Romania": "UEFA",
    "Scotland": "UEFA",
    "Serbia and Montenegro": "UEFA",
    "Slovenia": "UEFA",
    "Turkey": "UEFA",
    "Ukraine": "UEFA",
    "Wales": "UEFA",

    # CONMEBOL
    "Brazil": "CONMEBOL",
    "Argentina": "CONMEBOL",
    "Uruguay": "CONMEBOL",
    "Chile": "CONMEBOL",
    "Colombia": "CONMEBOL",
    "Peru": "CONMEBOL",
    "Ecuador": "CONMEBOL",
    "Paraguay": "CONMEBOL",
    "Bolivia": "CONMEBOL",
    "Venezuela": "CONMEBOL",

    # AFC
    "Japan": "AFC",
    "South Korea": "AFC",
    "Korea Republic": "AFC",
    "IR Iran": "AFC",
    "Iran": "AFC",
    "Saudi Arabia": "AFC",
    "Australia": "AFC",
    "North Korea": "AFC",
    "China": "AFC",
    "Qatar": "AFC",

    # CAF
    "Nigeria": "CAF",
    "Ghana": "CAF",
    "Cameroon": "CAF",
    "Senegal": "CAF",
    "Ivory Coast": "CAF",
    "Côte d’Ivoire": "CAF",
    "Morocco": "CAF",
    "Tunisia": "CAF",
    "Algeria": "CAF",
    "Egypt": "CAF",
    "Angola": "CAF",
    "South Africa": "CAF",
    "Togo": "CAF",

    # CONCACAF
    "United States": "CONCACAF",
    "Mexico": "CONCACAF",
    "Costa Rica": "CONCACAF",
    "Honduras": "CONCACAF",
    "Jamaica": "CONCACAF",
    "Panama": "CONCACAF",
    "Trinidad and Tobago": "CONCACAF",
    "Canada": "CONCACAF",

    # OFC
    "New Zealand": "OFC",
}

def norm(s):
    s = str(s).strip()
    s = re.sub(r"\s+", " ", s)
    return s

def populate_confederation():
    wb = load_workbook(EXCEL_PATH)
    ws = wb[SHEET_NAME]

    header_to_col = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(HEADER_ROW, c).value
        if v:
            header_to_col[str(v).strip()] = c

    if "team_name" not in header_to_col:
        raise ValueError("team_name column not found")

    team_col = header_to_col["team_name"]

    # Create confederation column if missing
    if "confederation" not in header_to_col:
        conf_col = ws.max_column + 1
        ws.cell(HEADER_ROW, conf_col).value = "confederation"
    else:
        conf_col = header_to_col["confederation"]

    missing = set()
    rows_written = 0

    for r in range(DATA_START_ROW, ws.max_row + 1):
        team = ws.cell(r, team_col).value
        if team is None:
            continue

        team_clean = norm(team)
        conf = TEAM_TO_CONFED.get(team_clean)

        if conf is None:
            missing.add(team_clean)
            ws.cell(r, conf_col).value = None
        else:
            ws.cell(r, conf_col).value = conf
            rows_written += 1

    wb.save(EXCEL_PATH)

    print(f"Done. Wrote confederation for {rows_written} rows.")
    if missing:
        print("Teams missing confederation mapping:")
        for t in sorted(missing):
            print(" -", t)

if __name__ == "__main__":
    populate_confederation()
